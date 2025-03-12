from langchain_community.vectorstores import FAISS
from langchain.schema import Document
from config import settings as config_settings
from docx import Document as DocxDocument
from utils.get_embeddings import embeddings
import os

def load_faiss_index(index_path):
    vectorstore = FAISS.load_local(
        index_path, embeddings, allow_dangerous_deserialization=True
    )
    return vectorstore


def get_common_issues_idx(paragraphs):
    for i in range(len(paragraphs)):
        if paragraphs[i].text.strip().lower() == "common issues":
            return i
    return -1


def create_faiss_index(documents, index_path):
    vectorstore = FAISS.from_documents(documents, embeddings)
    vectorstore.save_local(index_path)


def load_error_codes(docx_file):
    """Parses error codes from a structured DOCX file and creates documents with metadata."""
    doc = DocxDocument(docx_file)

    paragraphs = doc.paragraphs

    common_issues_idx = get_common_issues_idx(paragraphs)

    error_data = []
    current_error = {"Code": "", "Content": []}

    for i in range(1, len(paragraphs)):
        text = paragraphs[i].text.strip()
        style = paragraphs[i].style.name

        if i == common_issues_idx:
            continue

        if style.startswith("Heading") and text:
            if current_error["Code"]:
                error_data.append(current_error)

            current_error = {"Code": f"Code: {text}", "Content": []}
        elif (not style.startswith("Heading")) and text:
            current_error["Content"].append(text)

    if current_error["Code"]:
        error_data.append(current_error)

    documents = []
    for entry in error_data:
        documents.append(
            Document(
                page_content="Error Code: {0}\n\nContent: {1}".format(
                    entry["Code"], "\n".join(entry["Content"])
                ),
                metadata={
                    "Error Code": entry["Code"],
                },
            )
        )

    return documents


doc_file_path = config_settings.ERROR_CODE_DATA_SOURCE_PATH
indexes_path = config_settings.VECTORE_DATASTORE_PATH

# Only create vector store if it does not exist
if not os.path.exists(indexes_path):
    documents = load_error_codes(doc_file_path)
    create_faiss_index(documents, indexes_path)
    print("Vector store created successfully.")
else:
    print("Vector store already exists. Skipping creation.")

faiss_indexes = load_faiss_index(indexes_path)
